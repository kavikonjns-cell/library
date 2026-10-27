"""Exports: an Excel workbook of all data and a self-contained catalogue HTML file."""
from __future__ import annotations

import base64
import io
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import config, jalali, services
from .db import connect

HERE = Path(__file__).resolve().parent

LOAN_STATUS = {"reserved": "رزرو", "active": "در امانت", "returned": "بازگشته", "lost": "گم‌شده", "cancelled": "لغو"}
PAY_STATUS = {"paid": "پرداخت شد", "pending": "در انتظار", "rejected": "رد شد"}


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="9E2443")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in rows[:200]]) + 2
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(width, 8), 60)
    ws.freeze_panes = "A2"


def excel_bytes() -> bytes:
    j = jalali.to_jalali
    wb = Workbook()
    wb.remove(wb.active)
    with connect() as conn:
        books = conn.execute("SELECT * FROM books ORDER BY code").fetchall()
        rows = []
        for b in books:
            loan = conn.execute("SELECT l.*, m.name FROM loans l JOIN members m ON m.code = l.member_code"
                                " WHERE l.book_code = ? AND l.status IN ('reserved','active')", (b["code"],)).fetchone()
            times = conn.execute("SELECT COUNT(*) FROM loans WHERE book_code = ? AND status IN ('active','returned','lost')",
                                 (b["code"],)).fetchone()[0]
            state = "خارج از قفسه" if not b["active"] else (LOAN_STATUS[loan["status"]] if loan else "موجود")
            rows.append([b["code"], b["title"], b["author"], b["translator"], b["publisher"], b["year"], b["topic1"],
                         b["topic2"], b["short_intro"], b["condition"], b["source"], b["added_at"], state,
                         loan["name"] if loan else "", j(loan["due_date"] or loan["reserve_expires"]) if loan else "",
                         times, b["note"]])
        _sheet(wb, "کتاب‌ها", ["کد کتاب", "عنوان", "نویسنده / گردآورنده", "مترجم", "ناشر", "سال نشر", "موضوع اصلی",
                               "موضوع دوم", "معرفی کوتاه", "وضعیت فیزیکی", "منبع", "تاریخ ورود", "وضعیت فعلی", "در دستِ",
                               "تا تاریخ", "دفعات امانت", "یادداشت"], rows)

        members = conn.execute("SELECT * FROM members ORDER BY seq").fetchall()
        _sheet(wb, "اعضا", ["کد عضو", "نام و نام خانوادگی", "نقش", "توضیح نقش", "نام کودک", "گروه کودک", "شماره تماس",
                            "شماره تلگرام", "تاریخ عضویت یا تمدید", "اعتبار تا", "حق عضویت", "پرداخت", "وضعیت",
                            "موضوع‌های مورد علاقه", "پیشنهاد کتاب", "آی‌دی تلگرام", "یادداشت"],
               [[m["code"], m["name"], m["role"], m["role_other"], m["child_name"], m["child_group"], m["phone"],
                 m["tg_phone"], j(m["registered_at"]), j(m["valid_until"]), m["fee"], PAY_STATUS[m["payment_status"]],
                 "باطل‌شده" if m["status"] == "revoked" else "فعال", m["topics"], m["wish"],
                 ("@" + m["tg_username"]) if m["tg_username"] else (m["tg_user_id"] or ""), m["note"]] for m in members])

        loans = conn.execute("SELECT l.*, b.title, m.name, m.role FROM loans l JOIN books b ON b.code = l.book_code"
                             " JOIN members m ON m.code = l.member_code ORDER BY l.id").fetchall()
        today = jalali.today()
        _sheet(wb, "امانت‌ها", ["ردیف", "کد کتاب", "عنوان کتاب", "کد عضو", "نام عضو", "نقش", "تاریخ درخواست",
                                "مدت امانت (روز)", "تاریخ تحویل", "مهلت بازگشت", "تحویل‌دهنده", "تاریخ بازگشت",
                                "تحویل‌گیرنده", "وضعیت کتاب هنگام بازگشت", "وضعیت امانت", "روزهای تأخیر",
                                "جریمه (تومان)", "تسویه جریمه", "بازخورد خواننده"],
               [[l["id"], l["book_code"], l["title"], l["member_code"], l["name"], l["role"], j(l["requested_at"]),
                 l["requested_days"], j(l["handed_at"]), j(l["due_date"]), l["handed_by"], j(l["returned_at"]),
                 l["received_by"], l["return_condition"], LOAN_STATUS[l["status"]],
                 services.late_days(l, today) if l["status"] in ("active", "returned") else "",
                 l["fine"] or (services.fine_for(services.late_days(l, today)) if l["status"] == "active" else 0),
                 ("بله" if l["fine_paid"] else "خیر") if l["fine"] else "", l["feedback"]] for l in loans])

        waits = conn.execute("SELECT w.*, b.title, m.name FROM waitlist w JOIN books b ON b.code = w.base_code"
                             " JOIN members m ON m.code = w.member_code ORDER BY w.id").fetchall()
        wstate = {"waiting": "در انتظار", "fulfilled": "رزرو شد", "cancelled": "لغو"}
        _sheet(wb, "صف انتظار", ["کد کتاب", "عنوان", "کد عضو", "نام عضو", "تاریخ درخواست", "مدت درخواستی", "وضعیت"],
               [[w["base_code"], w["title"], w["member_code"], w["name"], j(w["requested_at"]), w["requested_days"],
                 wstate[w["status"]]] for w in waits])

        pays = conn.execute("SELECT * FROM payments ORDER BY date, id").fetchall()
        balance, prow = 0, []
        for p in pays:
            balance += p["amount"] if p["direction"] == "in" else -p["amount"]
            prow.append([p["id"], j(p["date"]), p["category"], "دریافت" if p["direction"] == "in" else "پرداخت",
                         p["description"], p["member_code"], p["book_code"], p["loan_id"] or "", p["method"], p["ref"],
                         p["amount"] if p["direction"] == "in" else "", p["amount"] if p["direction"] == "out" else "",
                         balance])
        _sheet(wb, "حسابداری", ["ردیف", "تاریخ", "دسته", "نوع", "شرح", "کد عضو", "کد کتاب", "شمارۀ امانت",
                                "روش پرداخت", "شمارۀ پیگیری", "دریافت", "پرداخت", "ماندۀ صندوق"], prow)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _data_uri(path: Path, mime: str) -> str:
    return f"data:{mime};base64," + base64.b64encode(path.read_bytes()).decode()


def static_catalog_html(groups) -> str:
    """The catalogue as one offline HTML file (images inlined), like the original book list.

    When opened with internet access it refreshes the shelf status from the live site.
    """
    from .web import templates  # late import: web imports this module

    covers = HERE / "static" / "covers"

    def cover(code: str) -> str:
        p = covers / f"{code}.jpg"
        return _data_uri(p, "image/jpeg") if p.exists() else ""

    css = (HERE / "static" / "style.css").read_text(encoding="utf-8")
    tpl = templates.get_template("catalog.html")
    return tpl.render(
        groups=groups, today=jalali.to_jalali(jalali.today()), static=True, inline_css=css,
        logo=_data_uri(HERE / "static" / "logo.png", "image/png"), cover_url=cover,
        day_choices=config.LOAN_DAY_CHOICES, default_days=config.LOAN_DEFAULT_DAYS,
    )
